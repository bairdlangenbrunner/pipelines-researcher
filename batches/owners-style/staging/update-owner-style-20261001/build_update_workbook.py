#!/usr/bin/env python3
"""Build the owner-style normalization workbook from staged_updates.json (per-batch update
builder, pattern of batches/united-states-oil/staging/update-delaware-express/).

Tabs (README first; no wrap anywhere — clip):
  OO_Names           THE decision surface: one line per distinct current spelling -> styled form,
                     with basis / entity id / dropped acronym / cells (Owner<N> + Operator) / PIDs.
                     Reject a NAME here and every one of its cells drops out of the paste.
  OO_OperatorsOwners paste-ready mirror of the "Pipeline operators/owners" tab (GID 1489950650):
                     the FULL column set in sheet order, current values prefilled for every touched
                     ProjectID, the re-spelled Operator / Owner<N> cells overlaid green (high), leading
                     SheetRow locator. Formula columns (AggregateOwners, Percentage Verification)
                     are left blank — never paste over the live formulas.
  OO_Changes         one line per cell: SheetRow, ProjectID, column, current, proposed, basis, evidence.
  OO_HeldBack        every cell the styler WOULD change but this slice does not stage, by reason —
                     the ruling surface for the next slice.

    python batches/owners-style/staging/update-owner-style-20261001/build_update_workbook.py \
        --output batches/owners-style/deliverables/pipelines_batch_<stamp>_owners-style_update.xlsx
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[3]
sys.path.insert(0, str(REPO / "scripts"))
from build_recon_workbook import CONF_FILL, HEADER_FILL, HEADER_FONT  # noqa: E402
from stage_owner_style import col_order  # noqa: E402

GREEN = CONF_FILL["green"]
CLIP = Alignment(wrap_text=False, vertical="top")
HDR_ALIGN = Alignment(wrap_text=False, vertical="center", horizontal="center")
FORMULA_COLS = {"AggregateOwners", "Percentage Verification"}
CELL_MAX = 32000


def _hdr(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(1, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HDR_ALIGN
    ws.freeze_panes = "A2"


def _clip_all(ws):
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = CLIP


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _s(v):
    v = "" if v is None else str(v)
    return v if len(v) <= CELL_MAX else v[:CELL_MAX - 1] + "…"


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--staging", type=Path, default=HERE)
    ap.add_argument("--output", required=True, type=Path)
    a = ap.parse_args()

    d = json.loads((a.staging / "staged_updates.json").read_text(encoding="utf-8"))
    meta, names, rows, held = d["meta"], d["names"], d["rows"], d["held_back"]
    oo_csv = REPO / "data" / meta["owners_csv"]
    with oo_csv.open(newline="", encoding="utf-8") as fh:
        raw = list(csv.reader(fh))
    hdr = [c.strip() for c in raw[1]]
    by_pid = {}
    for ri in range(2, len(raw)):
        r = raw[ri]
        if len(r) < len(hdr):
            r = r + [""] * (len(hdr) - len(r))
        pid = r[hdr.index("ProjectID")].strip()
        if pid:
            by_pid[pid] = (ri + 1, r)

    changed_names = {}
    for pid, row in rows.items():
        for col, ch in row["changes"].items():
            e = changed_names.setdefault(ch["old"], {"new": ch["new"], "basis": ch["basis"], "entity_id": ch["entity_id"],
                                                     "legal_form": ch["legal_form"], "aliases": ch["aliases"],
                                                     "flags": ch["flags"], "cells": 0, "operator_cells": 0,
                                                     "pids": set()})
            e["cells"] += 1
            if col == "Operator":
                e["operator_cells"] += 1
            e["pids"].add(pid)
    n_changes = sum(len(r["changes"]) for r in rows.values())
    counts = meta["counts"]
    orphans = {pid for pid in rows if by_pid[pid][1][hdr.index("PipelineName")].strip() == "#N/A"}

    wb = Workbook()

    # README
    ws = wb.active
    ws.title = "README"
    readme = [
        ["Owner-style normalization — existing Operator + Owner<N> cells, slice 1 (gazetteer-backed, high confidence)"],
        [""],
        ["Mode", "update (style-only; no research, no URL changes)"],
        ["Scope", meta["scope"]["description"]],
        ["Snapshot", f"data/{meta['owners_csv']} (operators/owners tab, GID 1489950650)"],
        ["Staged", f"batches/owners-style/staging/{meta['batch']}/staged_updates.json"],
        ["Date", meta["researched"]],
        ["Method", meta["method"]],
        ["Slice", meta["slice"]],
        ["URL verification", meta["url_verification"]],
        [""],
        ["WHAT THIS IS", "The ownership team writes an immediate owner as its full legal name plus a trailing short legal form, no punctuation, no trailing acronym (docs/reference/owner_style.md). The Operator column is written the same way (Baird 2026-10-01) — it is the same kind of entity name, and on most rows the operator IS Owner1. These cells already name the right entity; only the SPELLING moves to the team's form — e.g. 'Sonatrach' -> 'Sonatrach SpA', 'Iraq Ministry of Oil' -> 'Ministry of Oil (Iraq)', 'Saudi Aramco' -> 'Saudi Arabian Oil Co'. Every proposed form is either the gazetteer's own spelling for that entity (exact match) or reached through a CONFIRMED alias in data/owner_aliases.json. Nothing fuzzy, nothing inferred."],
        ["HOW TO REVIEW", "Decide by NAME on OO_Names (one line per current spelling). A rejected name removes all of its cells from the paste; an accepted one is pasted from OO_OperatorsOwners by SheetRow/ProjectID. Dropped acronyms ('… (SNGPL)') are kept as aliases in data/owner_aliases.json, not lost."],
        ["WHAT IS NOT HERE", f"{sum(counts['held_back'].values())} cells the styler would also change but this slice holds back ({counts.get('held_back_operator', 0)} of them Operator), listed on OO_HeldBack by reason: " + ", ".join(f"{k}={v}" for k, v in counts["held_back"].items()) + ". stem_medium (one gazetteer stem match, form adopted at medium) and rules_only (punctuation / long-form rules, no gazetteer hit) are the natural slice 2; alias_candidate / form_conflict / acronym_not_initials need a ruling; row_collision = two Owner<N> slots on one row would become one name, so the shares must be merged by hand; multi_operator = one Operator cell lists several operators ('A; B') — the sheet has no Operator1..N, so the cell's shape is a ruling before its spelling. Comma-joined operator lists ('TC Energy, IEnova') are not split (a comma is also legal-name punctuation) and sit under rules_only."],
        ["NOT TOUCHED", "Owner [ref] / Operator [ref] (refs are additive and none is proposed), Owner<N>% (shares unchanged), OperatorLocalLanguage, LastUpdated, Researcher. AggregateOwners and the tracker's Owner / Operator / Parent columns are formulas and follow on their own — never paste those columns."],
        [""],
        ["Headline", meta["summary"]],
        ["", f"{len(changed_names)} distinct names -> {n_changes} cells ({counts.get('changes_owner', n_changes)} Owner<N> + {counts.get('changes_operator', 0)} Operator) on {len(rows)} ProjectIDs; {counts['cells_in']} non-blank name cells / {counts['names_in']} distinct names in the snapshot."],
        [""],
        ["Color key"],
        ["green", "high confidence — the proposed spelling (exact gazetteer entry or confirmed alias). Only re-spelled Operator / Owner<N> cells are colored; everything else on the mirror is the current sheet value for context."],
        [""],
        ["Sheets"],
        ["OO_Names", "decision surface — one line per current spelling: proposed form, basis, GEM entity id, legal form, dropped acronym / aliases, cells (and how many of them are Operator), PIDs"],
        ["OO_OperatorsOwners", "paste surface — full operators/owners column set in sheet order for every touched ProjectID, re-spelled cells green, leading SheetRow; formula columns left blank"],
        ["OO_Changes", "one line per re-spelled cell with its evidence"],
        ["OO_HeldBack", "would-change cells NOT in this slice, by reason (next slice + rulings)"],
        [""],
        ["Known oddities to eyeball", "'Gaz-System' -> 'Operator Gazociagow Przesylowych GAZ-SYSTEM SP zoo', 'GAIL (India) Ltd' -> 'GAIL Ltd', 'Dongying United Petroleum Chemical Co Ltd' -> 'Dongying UNITED Petroleum Chemical Co Ltd' and 'Petroperú' -> 'Petroperu' are the gazetteer's own spellings for those entity ids — adopted because that is the team's record, flagged here because they look unusual. 'Saudi Aramco' and 'Aramco' both -> 'Saudi Arabian Oil Co'; 'SNTGN Transgaz SA' and 'Transgaz' both -> 'Transgaz SA'; 'AB Amber Grid' and 'Amber Grid' both -> 'Amber Grid AB' (same entity, two current spellings). On the Operator column 'SOCAR' -> 'State Oil Company of Azerbaijan Republic' and 'Pemex' -> 'Petróleos Mexicanos EPE' are confirmed aliases, so the short trade name gives way to the team's legal name."],
        ["Orphan rows", "The sheet's own lookup columns (PipelineName … PCI6) show #N/A on " + ", ".join(sorted(orphans)) + " — their ProjectID has no match in the tracker tabs. Carried as blank here; the Owner<N> re-spelling still applies, but the rows themselves are a data-health question (deleted tracker row? mistyped id?)."],
        ["Review app", "This is an update store (staged_updates.json); the review app reads deep-sweep stores only, so this batch is decided in the workbook."],
    ]
    for r in readme:
        ws.append([_s(x) for x in r])
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 180
    _clip_all(ws)
    ws["A1"].font = Font(bold=True, size=14)

    # OO_Names
    ws = wb.create_sheet("OO_Names")
    cols = ["Current spelling", "Proposed (team style)", "basis", "GEM entity id", "legal form",
            "dropped acronym / aliases", "flags", "cells", "of which Operator", "PIDs", "ProjectIDs"]
    ws.append(cols)
    for old, e in sorted(changed_names.items(), key=lambda kv: (-kv[1]["cells"], kv[0])):
        ws.append([_s(x) for x in [old, e["new"], e["basis"], e["entity_id"], e["legal_form"],
                                   " | ".join(e["aliases"]), " | ".join(e["flags"]), e["cells"],
                                   e["operator_cells"], len(e["pids"]), " ".join(sorted(e["pids"]))]])
        ws.cell(ws.max_row, 2).fill = GREEN
    _hdr(ws, len(cols))
    _widths(ws, [44, 44, 8, 16, 10, 36, 22, 7, 9, 6, 80])
    _clip_all(ws)

    # OO_OperatorsOwners (full mirror, touched PIDs)
    ws = wb.create_sheet("OO_OperatorsOwners")
    ws.append(["SheetRow"] + hdr)
    for pid in sorted(rows, key=lambda p: rows[p]["sheet_row"]):
        sheet_row, r = by_pid[pid]
        # the sheet's own lookup columns show #N/A on a row whose ProjectID has no tracker match;
        # carry that as blank (a literal "#N/A" string reads as an error cell) — listed in the README
        vals = [sheet_row] + [("" if (h in FORMULA_COLS or r[i].strip() == "#N/A") else _s(r[i]))
                              for i, h in enumerate(hdr)]
        for col, ch in rows[pid]["changes"].items():
            vals[1 + hdr.index(col)] = ch["new"]
        ws.append(vals)
        rn = ws.max_row
        for col in rows[pid]["changes"]:
            ws.cell(rn, 2 + hdr.index(col)).fill = GREEN
    _hdr(ws, len(hdr) + 1)
    w = [9] + [14] * len(hdr)
    for name, width in (("PipelineName", 34), ("SegmentName", 26), ("Wiki", 30), ("Notes/Links", 30),
                        ("Operator [ref]", 30), ("Owner [ref]", 30), ("Operator", 30)):
        if name in hdr:
            w[1 + hdr.index(name)] = width
    for i, h in enumerate(hdr):
        if h.startswith("Owner") and h[5:].isdigit():
            w[1 + i] = 34
    _widths(ws, w)
    _clip_all(ws)

    # OO_Changes
    ws = wb.create_sheet("OO_Changes")
    cols = ["SheetRow", "ProjectID", "PipelineName", "SegmentName", "Fuel", "Countries", "Column",
            "Current", "Proposed", "basis", "GEM entity id", "aliases", "evidence"]
    ws.append(cols)
    for pid in sorted(rows, key=lambda p: rows[p]["sheet_row"]):
        row = rows[pid]
        for col in sorted(row["changes"], key=col_order):
            ch = row["changes"][col]
            ws.append([_s(x) for x in [row["sheet_row"], pid, row["pipeline"], row["segment"], row["fuel"],
                                       row["countries"], col, ch["old"], ch["new"], ch["basis"],
                                       ch["entity_id"], " | ".join(ch["aliases"]), ch["evidence"]]])
            ws.cell(ws.max_row, 9).fill = GREEN
    _hdr(ws, len(cols))
    _widths(ws, [9, 10, 34, 26, 6, 18, 9, 40, 40, 8, 16, 24, 70])
    _clip_all(ws)

    # OO_HeldBack
    ws = wb.create_sheet("OO_HeldBack")
    cols = ["reason", "SheetRow", "ProjectID", "Column", "Current", "Styler would write", "basis", "confidence",
            "GEM entity id", "flags", "unadopted candidate", "note"]
    ws.append(cols)
    for h in held:
        ws.append([_s(x) for x in [h["reason"], h["sheet_row"], h["project_id"], h["column"], h["current"],
                                   h["styled"], h["basis"], h["confidence"], h["entity_id"],
                                   " | ".join(h["flags"]), h.get("candidate", ""), h.get("note", "")]])
    _hdr(ws, len(cols))
    _widths(ws, [20, 9, 10, 9, 44, 44, 8, 10, 16, 28, 40, 50])
    _clip_all(ws)

    a.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(a.output)
    print(f"wrote {a.output}  ({len(wb.sheetnames)} sheets: {', '.join(wb.sheetnames)})")
    print(f"  names={len(changed_names)} cells={n_changes} pids={len(rows)} held_back={len(held)}")
    print(f"  next: python scripts/recalc.py {a.output}")


if __name__ == "__main__":
    main()
