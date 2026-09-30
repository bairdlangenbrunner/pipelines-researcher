"""build_ref_workbook.py --decisions: paste surfaces carry only lines a person accepted in the
review app; detail tabs gain Decision/DecisionNote; rejects + suggests land on a Declined tab.
Without the flag the output is unchanged (all four proposals on the paste surface)."""
import json
import sys

import pytest
from openpyxl import load_workbook

import build_ref_workbook as brw
from review_fixture import SNAP_NAME, make_scope


def _dec(key, decision, note="", suggested="", reviewer="Tester"):
    return {"key": key, "dir": "d", "decision": decision, "note": note, "suggested_value": suggested,
            "reviewer": reviewer, "ts": "2099-01-01T00:00:00-05:00", "undecided": False}


DECISIONS = [
    _dec("d::P9001|4|Status [ref]", "accept"),                                  # ref line
    _dec("d::P9002|5|Capacity", "hold"),                                        # fill line
    _dec("d::P9005|9|Capacity", "reject", note="system aggregate, not the segment"),
    _dec("d::P9098|4|Capacity [ref]", "suggest", note="figure is for phase 2", suggested="8.0"),
]


def _scope(tmp_path, handoff=True, decisions=True):
    s = make_scope(tmp_path)
    deep = s["deep"]
    p = deep / "staged_resolutions.json"
    d = json.loads(p.read_text())
    d["meta"]["scope"]["csv"] = str(s["data"] / SNAP_NAME)       # the builder reads the snapshot
    p.write_text(json.dumps(d))
    if not handoff:
        (deep / "staged_actions.json").unlink()                  # -> the sweep Backend mirror
    if decisions:
        (deep / "review_decisions.json").write_text(
            json.dumps({"generated": "x", "decisions": {r["key"]: r for r in DECISIONS}}))
    return s


def _run(monkeypatch, staging, out, flag):
    argv = ["build_ref_workbook.py", "--staging", str(staging), "--output", str(out)]
    monkeypatch.setattr(sys, "argv", argv + (["--decisions"] if flag else []))
    brw.main()


def _rows(ws):
    h = [c.value for c in ws[1]]
    return h, [dict(zip(h, (c.value for c in row))) for row in ws.iter_rows(min_row=2)]


def _tinted(ws, pid, col):
    h = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if row[h.index("ProjectID")].value == pid:
            c = row[h.index(col)]
            return c.value, bool(c.fill and c.fill.fill_type)
    return None


def test_handoff_without_flag_all_four_on_paste_surface(tmp_path, monkeypatch):
    s = _scope(tmp_path)
    _run(monkeypatch, s["deep"], tmp_path / "out" / "t.xlsx", flag=False)
    wb = load_workbook(tmp_path / "out" / "t-actions.xlsx")
    ws = wb["Gas_AllFillsBackend"]
    assert {r["ProjectID"] for r in _rows(ws)[1]} == {"P9001", "P9002", "P9005", "P9098"}
    assert "http://new.example/b" in _tinted(ws, "P9001", "Status [ref]")[0]
    assert _tinted(ws, "P9002", "Capacity") == ("7.5", True)
    ev = load_workbook(tmp_path / "out" / "t-evidence.xlsx")
    assert "Gas_Declined" not in ev.sheetnames
    assert "Decision" not in _rows(ev["Gas_FillDetail"])[0]


def test_handoff_with_flag_only_accepts_paste(tmp_path, monkeypatch, capsys):
    s = _scope(tmp_path)
    _run(monkeypatch, s["deep"], tmp_path / "out" / "t.xlsx", flag=True)
    out = capsys.readouterr().out
    assert "decisions: accept=1 hold=1 reject=1 suggest=1" in out
    assert "1 of 1 staging dir(s) had a decision file" in out

    wb = load_workbook(tmp_path / "out" / "t-actions.xlsx")
    ws = wb["Gas_AllFillsBackend"]
    assert [r["ProjectID"] for r in _rows(ws)[1]] == ["P9001"]
    val, tint = _tinted(ws, "P9001", "Status [ref]")
    assert "http://new.example/b" in val and tint
    assert "Gas_OperatorsOwners" not in wb.sheetnames         # the owner line is undecided
    # undecided status change stays on the change list, with its Decision
    h, st = _rows(wb["Gas_StatusChanges"])
    assert [(r["ProjectID"], r["Decision"]) for r in st] == [("P9002", "undecided")]

    ev = load_workbook(tmp_path / "out" / "t-evidence.xlsx")
    h, fd = _rows(ev["Gas_FillDetail"])
    assert "Decision" in h and "DecisionNote" in h
    got = {r["ProjectID"]: r["Decision"] for r in fd}
    assert got["P9002"] == "hold" and got["P9005"] == "reject" and got["P9098"] == "suggest"
    assert "system aggregate" in next(r for r in fd if r["ProjectID"] == "P9005")["DecisionNote"]
    h, rd = _rows(ev["Gas_RefWorkDetail"])
    assert next(r for r in rd if r["ProjectID"] == "P9001"
                and r["Ref column"] == "Status [ref]")["Decision"] == "accept"

    h, dl = _rows(ev["Gas_Declined"])
    assert h == ["Kind", "ProjectID", "SheetRow", "Column", "Proposed", "Decision", "Reviewer",
                 "Note", "SuggestedValue", "Dir"]
    by = {r["ProjectID"]: r for r in dl}
    assert set(by) == {"P9005", "P9098"}
    assert by["P9005"]["Decision"] == "reject" and by["P9005"]["Note"] == "system aggregate, not the segment"
    assert by["P9098"]["Decision"] == "suggest" and by["P9098"]["SuggestedValue"] == "8.0"
    assert by["P9005"]["Kind"] == "fill" and by["P9005"]["Column"] == "Capacity"


def test_flag_without_decision_file_pastes_nothing(tmp_path, monkeypatch, capsys):
    s = _scope(tmp_path, decisions=False)
    _run(monkeypatch, s["deep"], tmp_path / "out" / "t.xlsx", flag=True)
    assert "0 of 1 staging dir(s) had a decision file" in capsys.readouterr().out
    wb = load_workbook(tmp_path / "out" / "t-actions.xlsx")
    assert "Gas_AllFillsBackend" not in wb.sheetnames


@pytest.mark.parametrize("flag", [False, True])
def test_sweep_backend_mirror(tmp_path, monkeypatch, flag):
    s = _scope(tmp_path, handoff=False)
    _run(monkeypatch, s["deep"], tmp_path / "out" / "t.xlsx", flag=flag)
    wb = load_workbook(tmp_path / "out" / "t.xlsx")
    ws = wb["Gas_Backend"]
    pids = {r["ProjectID"] for r in _rows(ws)[1]}
    assert {"P9001", "P9002", "P9005"} <= pids               # rows stay either way
    val, tint = _tinted(ws, "P9001", "Status [ref]")
    assert "http://new.example/b" in val and tint              # the accept pastes
    cap, cap_tint = _tinted(ws, "P9002", "Capacity")
    if flag:
        assert (cap or "") == "" and not cap_tint              # held fill: current (blank), no tint
        assert "Gas_Declined" in wb.sheetnames
        assert "Decision" in _rows(wb["Gas_Fills"])[0]
    else:
        assert cap == "7.5" and cap_tint
        assert "Gas_Declined" not in wb.sheetnames
